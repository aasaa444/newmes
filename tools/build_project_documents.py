from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "project-initiation"
OUTPUT = ROOT / "deliverables" / "project-initiation"

NAVY = "0F2B5B"
BLUE = "1E6FD9"
ACCENT = "F39800"
INK = "1F2A37"
GRAY = "5A6672"
LIGHT = "E8F1FC"
GRAYBG = "F2F4F7"
RED = "C8392E"


DELIVERABLES = [
    {
        "filename": "01-工业路由器装配MES-项目立项与需求分析书.docx",
        "title": "项目立项与需求分析书",
        "subtitle": "工业路由器整机装配 MES 试点候选项目",
        "doc_no": "MES-CHARTER-BRS-001",
        "sources": [
            "01-project-charter-and-scope.md",
            "02-business-requirements-specification.md",
            "06-evidence-and-discovery-plan.md",
        ],
    },
    {
        "filename": "02-工业路由器装配MES-解决方案蓝图.docx",
        "title": "MES 解决方案蓝图",
        "subtitle": "从单件追溯内核到企业集成与私有化交付",
        "doc_no": "MES-SOL-001",
        "sources": ["03-solution-blueprint.md", "06-evidence-and-discovery-plan.md"],
    },
    {
        "filename": "03-工业路由器装配MES-项目实施与验收方案.docx",
        "title": "项目实施与验收方案",
        "subtitle": "实施治理、UAT、上线准备与技术验证证据",
        "doc_no": "MES-IMP-TST-001",
        "sources": [
            "04-implementation-and-governance-plan.md",
            "05-test-uat-and-acceptance-plan.md",
        ],
    },
]


def set_run_font(run, name="Calibri", east_asia="Microsoft YaHei"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)
    for cell in row.cells:
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.keep_with_next = True


def keep_table_row_together(row):
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def set_table_widths(table, widths_dxa):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        keep_table_row_together(row)
        for index, (cell, width) in enumerate(zip(row.cells, widths_dxa)):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def set_cell_text(cell, text, bold=False, color=INK, size=9.5):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.05
    run = paragraph.add_run(text.strip())
    set_run_font(run)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def add_field(paragraph, field_name):
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field_name
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char, instr, fld_end])
    set_run_font(run)
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(GRAY)


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, NAVY, 8, 4),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.167


def configure_section(section, short_title):
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(short_title)
    set_run_font(run)
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(GRAY)
    p_pr = p._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "3")
    bottom.set(qn("w:color"), "D8DEE5")
    p_bdr.append(bottom)
    p_pr.append(p_bdr)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = fp.add_run("工业路由器装配 MES 试点候选 | 李先生 | ")
    set_run_font(run)
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(GRAY)
    add_field(fp, "PAGE")


def add_cover(doc, title, subtitle, doc_no):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(70)
    p.paragraph_format.space_after = Pt(12)
    run = p.add_run("MES PROJECT DELIVERY PACK")
    set_run_font(run)
    run.font.size = Pt(11)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(ACCENT)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    run = p.add_run(title)
    set_run_font(run)
    run.font.size = Pt(31)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(NAVY)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(28)
    run = p.add_run(subtitle)
    set_run_font(run)
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor.from_string(GRAY)

    table = doc.add_table(rows=4, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"
    values = [
        ("文档编号", doc_no),
        ("版本/日期", "0.1 / 2026-07-29"),
        ("编制人", "李先生"),
        ("可信度", "试点候选；尚未完成制造从业者评审或真实现场验证"),
    ]
    for row, (label, value) in zip(table.rows, values):
        set_cell_text(row.cells[0], label, bold=True, color=NAVY, size=10)
        shade(row.cells[0], LIGHT)
        set_cell_text(row.cells[1], value, size=10)
    set_table_widths(table, [2200, 7160])

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(26)
    p.paragraph_format.space_after = Pt(0)
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "FDF3E0")
    p_pr.append(shd)
    run = p.add_run("使用边界：本文件是公开证据驱动的候选方案，不代表已签约、已上线或已现场验收。")
    set_run_font(run)
    run.font.size = Pt(10.5)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string("9A6300")

    doc.add_page_break()


def add_toc(doc, source_paths):
    doc.add_heading("文档导航", level=1)
    p = doc.add_paragraph("以下目录由当前单一事实源生成；正式发布前应在 Word 中更新自动目录字段。")
    p.runs[0].font.color.rgb = RGBColor.from_string(GRAY)
    for path in source_paths:
        for line in path.read_text(encoding="utf-8").splitlines():
            if line.startswith("## "):
                p = doc.add_paragraph(style="List Bullet")
                add_inline(p, line[3:].strip())
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.0
                for run in p.runs:
                    run.font.size = Pt(9.5)
    doc.add_page_break()


INLINE_PATTERN = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")


def add_inline(paragraph, text):
    cursor = 0
    for match in INLINE_PATTERN.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor:match.start()])
            set_run_font(run)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.font.bold = True
        else:
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.color.rgb = RGBColor.from_string(NAVY)
        set_run_font(run, name=run.font.name or "Calibri")
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        set_run_font(run)


def parse_table(lines, start):
    rows = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        cells = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
        rows.append(cells)
        index += 1
    if len(rows) >= 2 and all(re.fullmatch(r":?-{3,}:?", cell) for cell in rows[1]):
        rows.pop(1)
    return rows, index


def add_markdown_table(doc, rows):
    if not rows:
        return
    cols = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"
    for r_index, row in enumerate(rows):
        for c_index in range(cols):
            value = row[c_index] if c_index < len(row) else ""
            set_cell_text(table.cell(r_index, c_index), value, bold=r_index == 0, color=NAVY if r_index == 0 else INK)
            if r_index == 0:
                shade(table.cell(r_index, c_index), GRAYBG)
    set_repeat_table_header(table.rows[0])
    max_lengths = [max(len(row[c]) if c < len(row) else 0 for row in rows) for c in range(cols)]
    weights = [max(8, min(length, 42)) for length in max_lengths]
    total = sum(weights)
    widths = [int(9360 * weight / total) for weight in weights]
    widths[-1] += 9360 - sum(widths)
    set_table_widths(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def create_numbering_instance(doc, start=1):
    numbering = doc.part.numbering_part._element
    existing = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    num_id = max(existing, default=0) + 1

    style_num_id = doc.styles["List Number"]._element.pPr.numPr.numId.val
    base_num = next(
        node for node in numbering.findall(qn("w:num"))
        if int(node.get(qn("w:numId"))) == int(style_num_id)
    )
    abstract_id = base_num.find(qn("w:abstractNumId")).get(qn("w:val"))

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract = OxmlElement("w:abstractNumId")
    abstract.set(qn("w:val"), abstract_id)
    num.append(abstract)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start_override = OxmlElement("w:startOverride")
    start_override.set(qn("w:val"), str(start))
    override.append(start_override)
    num.append(override)
    numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num)


def add_markdown(doc, path, include_title=False):
    lines = path.read_text(encoding="utf-8").splitlines()
    index = 0
    active_num_id = None
    while index < len(lines):
        raw = lines[index]
        line = raw.strip()
        if not line:
            active_num_id = None
            index += 1
            continue
        if line.startswith("# "):
            active_num_id = None
            if include_title:
                doc.add_heading(line[2:].strip(), level=1)
            index += 1
            continue
        if line.startswith("## "):
            active_num_id = None
            doc.add_heading(line[3:].strip(), level=1)
            index += 1
            continue
        if line.startswith("### "):
            active_num_id = None
            doc.add_heading(line[4:].strip(), level=2)
            index += 1
            continue
        if line.startswith("#### "):
            active_num_id = None
            doc.add_heading(line[5:].strip(), level=3)
            index += 1
            continue
        if line.startswith("|"):
            active_num_id = None
            rows, index = parse_table(lines, index)
            add_markdown_table(doc, rows)
            continue
        if line.startswith("- "):
            active_num_id = None
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, line[2:].strip())
            index += 1
            continue
        if re.match(r"^\d+\.\s", line):
            p = doc.add_paragraph(style="List Number")
            if active_num_id is None:
                start = int(re.match(r"^(\d+)\.", line).group(1))
                active_num_id = create_numbering_instance(doc, start=start)
            apply_numbering(p, active_num_id)
            add_inline(p, re.sub(r"^\d+\.\s+", "", line))
            index += 1
            continue
        if line.startswith("> "):
            active_num_id = None
            p = doc.add_paragraph()
            p_pr = p._p.get_or_add_pPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), LIGHT)
            p_pr.append(shd)
            add_inline(p, line[2:].strip())
            for run in p.runs:
                run.font.color.rgb = RGBColor.from_string(NAVY)
            index += 1
            continue
        paragraph_lines = [line]
        active_num_id = None
        index += 1
        while index < len(lines):
            candidate = lines[index].strip()
            if not candidate or candidate.startswith(("#", "|", "- ", "> ")) or re.match(r"^\d+\.\s", candidate):
                break
            paragraph_lines.append(candidate)
            index += 1
        p = doc.add_paragraph()
        add_inline(p, " ".join(paragraph_lines))


def build_deliverable(spec):
    doc = Document()
    configure_styles(doc)
    configure_section(doc.sections[0], spec["title"])
    doc.core_properties.author = "李先生"
    doc.core_properties.last_modified_by = "李先生"
    doc.core_properties.title = spec["title"]
    doc.core_properties.subject = spec["subtitle"]
    doc.core_properties.category = "MES 项目前期交付物"
    doc.core_properties.keywords = "MES,工业路由器,需求分析,解决方案,项目管理"

    add_cover(doc, spec["title"], spec["subtitle"], spec["doc_no"])
    source_paths = [SOURCE / filename for filename in spec["sources"]]
    add_toc(doc, source_paths)
    for source_index, source_path in enumerate(source_paths):
        if source_index:
            doc.add_page_break()
        add_markdown(doc, source_path)

    OUTPUT.mkdir(parents=True, exist_ok=True)
    target = OUTPUT / spec["filename"]
    doc.save(target)
    return target


if __name__ == "__main__":
    for deliverable in DELIVERABLES:
        print(build_deliverable(deliverable))
